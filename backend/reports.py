"""In-memory, filter-aware exports with spreadsheet formula protection."""

from io import BytesIO

import pandas as pd
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def spreadsheet_safe(df):
    df = df.copy()
    for col in df.select_dtypes(include=["object", "string"]).columns:
        df[col] = df[col].map(
            lambda x: (
                "'" + x
                if isinstance(x, str) and x.lstrip().startswith(("=", "+", "-", "@", "\t", "\r"))
                else x
            )
        )
    return df


def workbook(analytics, start=None, end=None, country=None):
    summary = analytics.overview(start, end, country)
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(
            [
                {"Metric": k.replace("_", " ").title(), "Value": summary[k]}
                for k in ["revenue", "refunds", "net_revenue", "orders", "customers", "aov", "units"]
            ]
        ).to_excel(writer, sheet_name="Executive summary", index=False)
        frames = {
            "Revenue trend": pd.DataFrame(summary["daily"]),
            "Markets": pd.DataFrame(summary["markets"]),
            "Top products": pd.DataFrame(summary["products"]),
            "Data quality": pd.DataFrame([analytics.audit]),
            "Provenance": pd.DataFrame(
                [
                    {
                        "Dataset": analytics.name,
                        "Currency": "GBP",
                        "Sales": "Completed product purchases; guest sales included",
                        "Refunds": "Reported separately from gross sales",
                        "Filters": f"{start or 'All'} to {end or 'All'}; {country or 'All markets'}",
                    }
                ]
            ),
        }
        # Customer profiles intentionally use the full dataset; label scope explicitly.
        if not analytics.rfm.empty:
            frames["Customers (all history)"] = analytics.rfm
        from src.visual_toolkit import get_tool_comparison_matrix

        frames["Visualization comparison"] = get_tool_comparison_matrix()
        for name, df in frames.items():
            spreadsheet_safe(df).to_excel(writer, sheet_name=name, index=False)
        for ws in writer.book:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="18362F")
                cell.font = Font(color="FFFFFF", bold=True)
            for idx, column in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in list(column)[:100]) + 3
                ws.column_dimensions[get_column_letter(idx)].width = min(50, max(14, width))
        for sheet, chart_type, title in [
            ("Revenue trend", LineChart, "Daily sales (GBP)"),
            ("Markets", DoughnutChart, "Revenue by market"),
            ("Top products", BarChart, "Top products by revenue"),
        ]:
            ws = writer.book[sheet]
            if ws.max_row < 2:
                continue
            chart = chart_type()
            chart.title = title
            col = 3 if sheet == "Top products" else 2
            chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=ws.max_row), titles_from_data=True)
            chart.set_categories(
                Reference(ws, min_col=2 if sheet == "Top products" else 1, min_row=2, max_row=ws.max_row)
            )
            chart.width = 24
            chart.height = 12
            ws.add_chart(chart, "H2")
    return output.getvalue()
