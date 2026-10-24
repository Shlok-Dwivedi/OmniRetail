from io import BytesIO
import time

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from backend.main import create_app


def test_overview_filters_reconcile(client):
    full = client.get("/api/overview").json()
    uk = client.get("/api/overview?country=United%20Kingdom").json()
    france = client.get("/api/overview?country=France").json()
    assert full["revenue"] == uk["revenue"] + france["revenue"]
    assert abs(sum(d["revenue"] for d in full["daily"]) - full["revenue"]) < 0.01
    assert abs(full["aov"] * full["orders"] - full["revenue"]) <= full["orders"] * 0.0051
    empty = client.get("/api/overview?country=Missing").json()
    assert empty["revenue"] == 0 and empty["products"] == []


def test_customer_search_sort_and_pagination(client):
    result = client.get("/api/customers?page_size=5").json()
    assert len(result["items"]) == 5
    assert result["total"] == 48
    spend = [c["MonetaryGBP"] for c in result["items"]]
    assert spend == sorted(spend, reverse=True)
    search = client.get("/api/customers?search=France").json()
    assert all(c["Country"] == "France" for c in search["items"])
    assert client.get("/api/customers?page=-1").status_code == 422
    assert client.get("/api/customers?sort=not_a_column").status_code == 422
    assert client.get("/api/customers/summary").json()["total"] == 48


def test_invalid_dates_predictions_and_routes(client):
    for url in [
        "/api/overview?start=2026-05-01&end=2025-01-01",
        "/api/overview?start=invalid",
        "/api/forecast?horizon=1000",
    ]:
        assert client.get(url).status_code == 422
    assert client.post("/api/predict", json={"recency": 500, "tenure": 5}).status_code == 422
    assert client.post("/api/predict", json={"frequency": -1}).status_code == 422
    assert client.get("/api/import/nonexistent").status_code == 404
    assert client.get("/api/unknown").status_code == 404


def test_held_out_models_and_prediction(client):
    response = client.get("/api/models")
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["horizon"] == 90 and data["test_customers"] > 0
    assert data["train_customers"] + data["test_customers"] == 48
    assert "201" not in data["cutoff"]  # fixture remains independent of bundled UCI data
    prediction = client.post(
        "/api/predict", json={"recency": 20, "frequency": 10, "tenure": 120, "average_order_value": 50}
    )
    assert prediction.status_code == 200, prediction.text
    p = prediction.json()
    assert 0 <= p["lower"] <= p["estimate"] <= p["upper"]


def test_forecast_export_and_cache(client):
    result = client.get("/api/forecast?horizon=7")
    assert result.status_code == 200, result.text
    forecast = result.json()
    assert len(forecast["forecast"]) == 7
    assert forecast["rmse"] >= 0 and forecast["baseline_rmse"] >= 0
    for row in forecast["forecast"]:
        assert 0 <= row["LowerBoundGBP"] <= row["PredictedRevenueGBP"] <= row["UpperBoundGBP"]
    assert client.get("/api/forecast?horizon=7").json() == forecast
    csv = client.get("/api/export/forecast?horizon=7")
    assert csv.status_code == 200
    assert "PredictedRevenueGBP" in csv.text


def test_workbook_contains_charts_and_matches_filter(client):
    response = client.get("/api/export/workbook?country=France")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    expected = client.get("/api/overview?country=France").json()["revenue"]
    assert book["Executive summary"]["B2"].value == expected
    assert len(book["Revenue trend"]._charts) == 1
    assert len(book["Markets"]._charts) == 1
    assert "Data quality" in book.sheetnames
    assert "GBP" in book["Provenance"]["B2"].value


def wait_for_job(client, job_id):
    for _ in range(100):
        job = client.get(f"/api/import/{job_id}").json()
        if job["status"] != "processing":
            return job
        time.sleep(0.03)
    raise AssertionError("Import did not finish")


def test_import_is_persistent_and_failed_import_preserves_dataset(client, data_dir, raw_transactions):
    original = client.get("/api/meta").json()["version"]
    bad = client.post("/api/import", files={"file": ("broken.csv", b"wrong\n1\n", "text/csv")})
    assert bad.status_code == 202
    job = wait_for_job(client, bad.json()["id"])
    assert job["status"] == "failed" and "Missing columns" in job["message"]
    assert client.get("/api/meta").json()["version"] == original
    content = raw_transactions.head(10).to_csv(index=False).encode()
    imported = client.post("/api/import", files={"file": ("new-sales.csv", content, "text/csv")})
    assert imported.status_code == 202
    assert wait_for_job(client, imported.json()["id"])["status"] == "complete"
    version = client.get("/api/meta").json()["version"]
    assert version != original
    with TestClient(create_app(data_dir)) as restored:
        assert restored.get("/api/meta").json()["version"] == version
        assert restored.get("/api/overview").json()["rows"] == 10
        assert restored.get("/api/forecast").status_code == 422
    assert client.post("/api/import", files={"file": ("bad.exe", b"1")}).status_code == 422
    assert client.post("/api/import", files={"file": ("empty.csv", b"")}).status_code == 422


def test_exchange_rate_failure_is_explicit(client, monkeypatch):
    import requests

    def fail(*args, **kwargs):
        raise requests.Timeout("timeout")

    monkeypatch.setattr(requests.Session, "get", fail)
    result = client.post("/api/currency/refresh")
    assert result.status_code == 503
    assert client.get("/api/meta").json()["currency"]["rates"] == {"GBP": 1.0}


def test_security_headers(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.headers["x-content-type-options"] == "nosniff"


def test_out_of_range_dates_are_validation_errors(client):
    for value in ["0001-01-01", "9999-12-31"]:
        assert client.get("/api/overview", params={"start": value}).status_code == 422
